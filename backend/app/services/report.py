from __future__ import annotations

import math
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook

from ..errors import invalid_param, missing_required_param

TEMPLATE_CODES = ("overview", "group_summary", "top_n")


def _j(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        f = float(value)
        return None if math.isnan(f) else f
    if isinstance(value, (np.bool_,)):
        return bool(value)
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (pd.Timestamp, np.datetime64)):
        return str(value)
    if pd.isna(value):
        return None
    return value


def _is_numeric(df: pd.DataFrame, column: str) -> bool:
    return pd.api.types.is_numeric_dtype(df[column])


def _envelope(df: pd.DataFrame, file_meta: dict) -> dict:
    return {
        "file": file_meta,
        "dataset": {"row_count": int(len(df)), "column_count": int(len(df.columns))},
        "sections": [],
    }


def build_report_content(
    code: str, df: pd.DataFrame, file_meta: dict, params: dict
) -> dict:
    if code == "overview":
        content = _envelope(df, file_meta)
        content["sections"] = [_overview_fields(df), _overview_stats(df)]
        return content
    if code == "group_summary":
        content = _envelope(df, file_meta)
        content["sections"] = [_group_summary(df, params)]
        return content
    if code == "top_n":
        content = _envelope(df, file_meta)
        content["sections"] = [_top_n(df, params)]
        return content
    raise invalid_param(f"未知模板：{code}")


def _overview_fields(df: pd.DataFrame) -> dict:
    from .parser import infer_columns

    rows = [[c["name"], c["type"]] for c in infer_columns(df)]
    return {"type": "table", "title": "字段清单", "columns": ["字段", "类型"], "rows": rows}


def _overview_stats(df: pd.DataFrame) -> dict:
    items = []
    for name in df.columns:
        if _is_numeric(df, name):
            series = df[name]
            items.append(
                {
                    "column": str(name),
                    "min": _j(series.min()),
                    "max": _j(series.max()),
                    "mean": _j(series.mean()),
                    "null_count": int(series.isna().sum()),
                }
            )
    return {"type": "stats", "title": "数值列统计", "items": items}


def _group_summary(df: pd.DataFrame, params: dict) -> dict:
    group_by = params.get("group_by")
    if not group_by:
        raise missing_required_param("缺少必选参数：分组列 group_by（FR-015）")
    if group_by not in df.columns:
        raise invalid_param(f"分组列不存在：{group_by}")

    agg = params.get("agg") or "sum"
    if agg not in ("sum", "mean"):
        raise invalid_param(f"聚合方式仅支持 sum/mean：{agg}")

    value_by = params.get("value_by")
    if value_by is not None:
        if value_by not in df.columns:
            raise invalid_param(f"数值列不存在：{value_by}")
        if not _is_numeric(df, value_by):
            raise invalid_param(f"列 {value_by} 不是数值列，无法聚合")

    grouped = df.groupby(group_by, dropna=False)
    counts = grouped.size().reset_index(name="行数")
    counts[group_by] = counts[group_by].apply(
        lambda v: "(空)" if pd.isna(v) else _j(v)
    )

    columns = [group_by, "行数"]
    if value_by:
        agg_series = grouped[value_by].agg(agg).reset_index(name=f"{value_by}_{agg}")
        agg_series[group_by] = agg_series[group_by].apply(
            lambda v: "(空)" if pd.isna(v) else _j(v)
        )
        counts = counts.merge(agg_series, on=group_by)
        columns.append(f"{value_by}_{agg}")

    counts = counts.sort_values("行数", ascending=False)
    rows = [[_j(v) for v in row] for row in counts[columns].values.tolist()]
    return {
        "type": "table",
        "title": f"分组汇总（按 {group_by}）",
        "columns": columns,
        "rows": rows,
    }


def _top_n(df: pd.DataFrame, params: dict) -> dict:
    sort_by = params.get("sort_by")
    if not sort_by:
        raise missing_required_param("缺少必选参数：排序列 sort_by（FR-015）")
    if sort_by not in df.columns:
        raise invalid_param(f"排序列不存在：{sort_by}")
    if not _is_numeric(df, sort_by):
        raise invalid_param(f"列 {sort_by} 不是数值列，无法排序")

    n = params.get("n")
    if n is None:
        n = 10
    try:
        n = int(n)
    except (TypeError, ValueError):
        raise invalid_param(f"N 必须是整数：{n}")
    if not 1 <= n <= 100:
        raise invalid_param(f"N 必须在 1-100 之间：{n}")

    top = df.sort_values(sort_by, ascending=False).head(n)
    columns = ["排名"] + [str(c) for c in df.columns]
    rows = []
    for rank, (_, row) in enumerate(top.iterrows(), start=1):
        rows.append([rank] + [_j(row[c]) for c in df.columns])
    return {
        "type": "table",
        "title": f"TOP {n} 排行（按 {sort_by} 降序）",
        "columns": columns,
        "rows": rows,
    }


def export_xlsx(content: dict, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    info = wb.create_sheet("报表信息")
    info.append(["文件", content["file"].get("name")])
    info.append(["生成时间", content["file"].get("generated_at")])
    info.append(["总行数", content["dataset"]["row_count"]])
    info.append(["总列数", content["dataset"]["column_count"]])

    for index, section in enumerate(content["sections"], start=1):
        title = f"{index}-{section['title']}"[:31]
        ws = wb.create_sheet(title)
        if section.get("type") == "table":
            ws.append(section["columns"])
            for row in section["rows"]:
                ws.append(row)
        elif section.get("type") == "stats":
            if section["items"]:
                keys = list(section["items"][0].keys())
                ws.append(keys)
                for item in section["items"]:
                    ws.append([item.get(k) for k in keys])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
